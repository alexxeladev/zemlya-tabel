"""Вахта: режим половины месяца и налог на официальную часть (task_vahta_taxes).

Налог = официальная выплата × ставка из настроек вахты; база разнесения по
юрлицам = «итого начислено» + налог. Проверочный пример — Караулов (вымышленный
охранник тех же фикстур, что в `test_vahta.py`): начислено 75 230, официальная
выплата 25 230, ставка 40 % → налог 10 092, база 85 322, расклад 35/60/5 →
29 862,70 / 51 193,20 / 4 266,10.

Налоги только в вахте: строки прочих подразделений не меняются ни на рубль —
это держит `TestOtherDepartmentsUntouched`.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.company_shares import EmployeeCompanyShare
from app.models.employees import Employee
from app.models.guard_settings import GuardSettings
from app.models.production_calendars import ProductionCalendar
from app.models.schedules import Schedule
from app.models.timesheet_entries import TimesheetEntry
from app.services.guard_duty import (
    create_assignment,
    employer_tax_percent,
    set_employer_tax_percent,
)
from app.services.guard_export import generate_guard_timesheet_excel
from app.services.guard_month import build_guard_month
from app.services.guard_payroll import calculate_guard_row, employer_tax
from app.services.payroll_statement import build_payroll_statement
from tests.test_vahta import (  # noqa: F401 — фикстуры модуля вахты
    FIRST_HALF,
    MONTH,
    YEAR,
    _auth,
    companies,
    crew,
    gbr_place,
    guard_dept,
    other_dept,
    rodionov,
    users,
    zone,
)

_ZERO = Decimal("0")
FORTY = Decimal("40")


def _days(first: int, last: int) -> set[int]:
    return set(range(first, last + 1))


def _karaulov(db, place, person, *, official_h1="25230", official_h2="0",
              days=FIRST_HALF, premium_h1="230"):
    """Проверочная строка: 15 смен × 5 000 + премия 230, оф. выплата 25 230."""
    assignment = create_assignment(
        db, year=YEAR, month=MONTH, place=place,
        position=person.primary_position, days=days,
    )
    assignment.premium_h1 = Decimal(premium_h1)
    assignment.official_payout_h1 = Decimal(official_h1)
    assignment.official_payout_h2 = Decimal(official_h2)
    db.commit()
    return assignment


def _admin(db) -> Employee:
    admin = Employee(full_name="Админ налогов", email="taxadmin@example.com",
                     role="admin", is_active=True, is_system_admin=True)
    db.add(admin)
    db.commit()
    return admin


# ── Чистый расчёт ─────────────────────────────────────────────────────────────

class TestTaxCalculation:
    def test_tax_is_official_payout_times_rate(self):
        assert employer_tax(Decimal("25230"), FORTY) == Decimal("10092.00")

    def test_tax_keeps_kopecks(self):
        """Decimal до копейки: 12 345,67 × 40 % = 4 938,268 → 4 938,27."""
        assert employer_tax(Decimal("12345.67"), FORTY) == Decimal("4938.27")

    def test_no_official_payout_no_tax(self):
        assert employer_tax(_ZERO, FORTY) == _ZERO

    def test_rate_is_a_parameter_not_a_constant(self):
        """Ставка приходит из настроек: при 30 % налог другой."""
        assert employer_tax(Decimal("25230"), Decimal("30")) == Decimal("7569.00")

    def test_karaulov_example(self):
        """Обязательный пример задачи: 75 230 + 10 092 = 85 322 → 35/60/5."""
        from app.services.guard_payroll import distribute_guard_amount

        row = calculate_guard_row(
            pay_type="per_shift", rate=Decimal("5000"), year=YEAR, month=MONTH,
            days=_days(1, 15), premium={1: Decimal("230")},
            official={1: Decimal("25230")}, tax_percent=FORTY,
        )
        assert row.accrued == Decimal("75230")
        assert row.official_payout == Decimal("25230")
        assert row.tax == Decimal("10092.00")
        assert row.distribution_base == Decimal("85322.00")

        amounts = distribute_guard_amount(
            row.distribution_base, {1: Decimal("35"), 2: Decimal("60"), 3: Decimal("5")}
        )
        assert amounts == {
            1: Decimal("29862.70"), 2: Decimal("51193.20"), 3: Decimal("4266.10"),
        }
        assert sum(amounts.values()) == Decimal("85322.00")

    def test_tax_does_not_touch_accrued_or_payout(self):
        """Налог — затрата компании: «итого начислено» и «к выплате» те же."""
        kwargs = dict(
            pay_type="per_shift", rate=Decimal("5000"), year=YEAR, month=MONTH,
            days=_days(1, 15), premium={1: Decimal("230")},
            official={1: Decimal("25230")},
        )
        with_tax = calculate_guard_row(**kwargs, tax_percent=FORTY)
        without = calculate_guard_row(**kwargs, tax_percent=_ZERO)
        assert with_tax.accrued == without.accrued == Decimal("75230")
        assert with_tax.net_payout == without.net_payout == Decimal("50000")

    def test_without_official_payout_base_equals_accrued(self):
        row = calculate_guard_row(
            pay_type="per_shift", rate=Decimal("5000"), year=YEAR, month=MONTH,
            days=_days(1, 15), premium={1: Decimal("230")}, tax_percent=FORTY,
        )
        assert row.tax == _ZERO
        assert row.distribution_base == row.accrued == Decimal("75230")


class TestTaxByHalves:
    def test_tax_is_computed_per_half(self):
        """Официальная выплата своя у каждой половины — налог тоже."""
        row = calculate_guard_row(
            pay_type="per_shift", rate=Decimal("5000"), year=YEAR, month=MONTH,
            days=_days(1, 31), premium={1: Decimal("230")},
            official={1: Decimal("25230"), 2: Decimal("10000")}, tax_percent=FORTY,
        )
        assert row.halves[1].tax == Decimal("10092.00")
        assert row.halves[2].tax == Decimal("4000.00")
        assert row.tax == Decimal("14092.00")
        assert row.halves[2].distribution_base == Decimal("84000.00")  # 16 × 5000 + 4000

    def test_only_half_keeps_its_own_sums(self):
        row = calculate_guard_row(
            pay_type="per_shift", rate=Decimal("5000"), year=YEAR, month=MONTH,
            days=_days(1, 31), premium={1: Decimal("230")},
            penalty={2: Decimal("500")},
            official={1: Decimal("25230"), 2: Decimal("10000")}, tax_percent=FORTY,
        )
        second = row.only_half(2)
        assert second.shifts == 16
        assert second.salary == Decimal("80000")
        assert second.premium == _ZERO
        assert second.penalty == Decimal("500")
        assert second.official_payout == Decimal("10000")
        assert second.accrued == Decimal("79500")
        assert second.net_payout == Decimal("69500")
        assert second.tax == Decimal("4000.00")
        assert second.distribution_base == Decimal("83500.00")


# ── Настройка ставки ──────────────────────────────────────────────────────────

class TestTaxRateSetting:
    def test_default_is_forty_percent_without_settings_row(self, db_session):
        assert db_session.query(GuardSettings).count() == 0
        assert employer_tax_percent(db_session) == Decimal("40")

    def test_rate_is_stored_and_used(self, db_session):
        set_employer_tax_percent(db_session, Decimal("30"))
        db_session.commit()
        assert employer_tax_percent(db_session) == Decimal("30")

    def test_api_admin_reads_and_changes_rate(self, client, users):
        headers = _auth(client, "admin")
        resp = client.get("/api/vahta/settings", headers=headers)
        assert resp.status_code == 200
        assert Decimal(resp.json()["employer_tax_percent"]) == Decimal("40")

        resp = client.patch(
            "/api/vahta/settings", json={"employer_tax_percent": "30"}, headers=headers
        )
        assert resp.status_code == 200
        assert Decimal(resp.json()["employer_tax_percent"]) == Decimal("30")
        assert Decimal(
            client.get("/api/vahta/settings", headers=headers).json()["employer_tax_percent"]
        ) == Decimal("30")

    def test_api_manager_can_change_rate(self, client, users):
        resp = client.patch(
            "/api/vahta/settings", json={"employer_tax_percent": "35"},
            headers=_auth(client, "manager"),
        )
        assert resp.status_code == 200

    def test_api_rejects_out_of_range(self, client, users):
        headers = _auth(client, "admin")
        for bad in ("-1", "101"):
            resp = client.patch(
                "/api/vahta/settings", json={"employer_tax_percent": bad}, headers=headers
            )
            assert resp.status_code == 422

    def test_api_accountant_cannot_change_rate(self, client, users):
        resp = client.patch(
            "/api/vahta/settings", json={"employer_tax_percent": "30"},
            headers=_auth(client, "accountant"),
        )
        assert resp.status_code == 403

    def test_api_timekeeper_does_not_see_rate(self, client, users):
        headers = _auth(client, "timekeeper")
        assert client.get("/api/vahta/settings", headers=headers).status_code == 403
        assert client.patch(
            "/api/vahta/settings", json={"employer_tax_percent": "30"}, headers=headers
        ).status_code == 403

    def test_changed_rate_changes_distribution(self, db_session, gbr_place, rodionov):
        _karaulov(db_session, gbr_place, rodionov)
        set_employer_tax_percent(db_session, Decimal("30"))
        db_session.commit()
        row = build_guard_month(db_session, _admin(db_session), YEAR, MONTH) \
            .zones[0].cards[0].rows[0]
        assert row.tax == Decimal("7569.00")
        assert row.distribution_base == Decimal("82799.00")


# ── Экран месяца ──────────────────────────────────────────────────────────────

class TestMonthScreen:
    def test_karaulov_row_on_screen(self, db_session, gbr_place, rodionov, companies):
        _karaulov(db_session, gbr_place, rodionov)
        month = build_guard_month(db_session, _admin(db_session), YEAR, MONTH)
        row = month.zones[0].cards[0].rows[0]
        assert row.accrued == Decimal("75230")
        assert row.tax == Decimal("10092.00")
        assert row.distribution_base == Decimal("85322.00")
        assert row.distribution == {
            companies["ZMO"].id: Decimal("29862.70"),
            companies["EKS"].id: Decimal("51193.20"),
            companies["SEC"].id: Decimal("4266.10"),
        }
        # Разница с «итого начислено» объяснена отдельными итогами.
        assert month.total_accrued == Decimal("75230")
        assert month.total_tax == Decimal("10092.00")
        assert month.total_distribution == Decimal("85322.00")
        assert month.total_distribution_base == Decimal("85322.00")
        assert month.employer_tax_percent == Decimal("40")
        assert sum(t.amount for t in month.company_totals) == Decimal("85322.00")

    def test_row_without_official_payout_is_unchanged(
        self, db_session, gbr_place, rodionov, companies
    ):
        _karaulov(db_session, gbr_place, rodionov, official_h1="0")
        row = build_guard_month(db_session, _admin(db_session), YEAR, MONTH) \
            .zones[0].cards[0].rows[0]
        assert row.tax == _ZERO
        assert row.distribution_base == row.accrued == Decimal("75230")
        # Прежние цифры task_vahta: 26 330,50 / 45 138,00 / 3 761,50.
        assert row.distribution == {
            companies["ZMO"].id: Decimal("26330.50"),
            companies["EKS"].id: Decimal("45138.00"),
            companies["SEC"].id: Decimal("3761.50"),
        }

    def test_timekeeper_gets_no_tax(self, client, users, db_session, gbr_place, rodionov):
        _karaulov(db_session, gbr_place, rodionov)
        data = client.get(
            f"/api/vahta/{YEAR}/{MONTH}", headers=_auth(client, "timekeeper")
        ).json()
        row = data["zones"][0]["cards"][0]["rows"][0]
        assert row["tax"] is None and row["distribution_base"] is None
        assert data["total_tax"] is None
        assert data["employer_tax_percent"] is None
        assert "10092" not in str(data)


class TestHalfView:
    """Переключатель: месяц целиком, 1–15, 16–конец."""

    def _setup(self, db, place, person):
        assignment = _karaulov(
            db, place, person, days=_days(1, 31), official_h2="10000",
        )
        assignment.penalty_h2 = Decimal("500")
        db.commit()

    def test_month_view(self, db_session, gbr_place, rodionov):
        self._setup(db_session, gbr_place, rodionov)
        month = build_guard_month(db_session, _admin(db_session), YEAR, MONTH)
        row = month.zones[0].cards[0].rows[0]
        assert month.view_half is None
        assert (month.first_day, month.last_day) == (1, 31)
        assert row.shifts == 31
        assert row.accrued == Decimal("154730")
        assert row.tax == Decimal("14092.00")
        assert [h.half for h in month.halves] == [1, 2]

    def test_first_half_view(self, db_session, gbr_place, rodionov, companies):
        self._setup(db_session, gbr_place, rodionov)
        month = build_guard_month(db_session, _admin(db_session), YEAR, MONTH, half=1)
        row = month.zones[0].cards[0].rows[0]
        assert month.view_half == 1
        assert (month.first_day, month.last_day) == (1, 15)
        assert row.shifts == 15
        assert row.salary == Decimal("75000")
        assert row.premium == Decimal("230")
        assert row.penalty == _ZERO
        assert row.official_payout == Decimal("25230")
        assert row.accrued == Decimal("75230")
        assert row.net_payout == Decimal("50000")
        assert row.tax == Decimal("10092.00")
        assert row.distribution[companies["ZMO"].id] == Decimal("29862.70")
        assert [h.half for h in row.halves] == [1]
        assert month.total_shifts == 15
        assert month.total_accrued == Decimal("75230")
        assert month.total_net_payout == Decimal("50000")
        assert sum(t.amount for t in month.company_totals) == Decimal("85322.00")
        assert [h.half for h in month.halves] == [1]
        assert month.zones[0].total_shifts == 15
        assert month.zones[0].total_accrued == Decimal("75230")

    def test_second_half_view(self, db_session, gbr_place, rodionov):
        self._setup(db_session, gbr_place, rodionov)
        month = build_guard_month(db_session, _admin(db_session), YEAR, MONTH, half=2)
        row = month.zones[0].cards[0].rows[0]
        assert (month.first_day, month.last_day) == (16, 31)
        assert row.shifts == 16
        assert row.salary == Decimal("80000")
        assert row.premium == _ZERO
        assert row.penalty == Decimal("500")
        assert row.official_payout == Decimal("10000")
        assert row.accrued == Decimal("79500")
        assert row.net_payout == Decimal("69500")
        # Налог — от официальной выплаты ЭТОЙ половины.
        assert row.tax == Decimal("4000.00")
        assert row.distribution_base == Decimal("83500.00")
        assert month.total_tax == Decimal("4000.00")
        assert sum(t.amount for t in month.company_totals) == Decimal("83500.00")

    def test_half_view_keeps_all_day_marks(self, db_session, gbr_place, rodionov):
        """Отметки — данные, а не итог: фронт шлёт набор дней строки целиком."""
        self._setup(db_session, gbr_place, rodionov)
        month = build_guard_month(db_session, _admin(db_session), YEAR, MONTH, half=2)
        assert month.zones[0].cards[0].rows[0].days == list(range(1, 32))

    def test_half_views_add_up_to_month(self, db_session, gbr_place, rodionov):
        self._setup(db_session, gbr_place, rodionov)
        admin = _admin(db_session)
        whole = build_guard_month(db_session, admin, YEAR, MONTH)
        first = build_guard_month(db_session, admin, YEAR, MONTH, half=1)
        second = build_guard_month(db_session, admin, YEAR, MONTH, half=2)
        for field in ("total_shifts", "total_accrued", "total_net_payout", "total_tax"):
            assert getattr(first, field) + getattr(second, field) == getattr(whole, field)

    def test_api_half_param(self, client, users, db_session, gbr_place, rodionov):
        self._setup(db_session, gbr_place, rodionov)
        headers = _auth(client, "admin")
        data = client.get(f"/api/vahta/{YEAR}/{MONTH}?half=2", headers=headers).json()
        assert data["view_half"] == 2
        assert data["first_day"] == 16
        assert data["zones"][0]["cards"][0]["rows"][0]["shifts"] == 16
        assert client.get(
            f"/api/vahta/{YEAR}/{MONTH}?half=3", headers=headers
        ).status_code == 422


# ── Общая ведомость и Excel ───────────────────────────────────────────────────

class TestStatement:
    def test_guard_row_distribution_includes_tax(
        self, db_session, gbr_place, rodionov, companies
    ):
        _karaulov(db_session, gbr_place, rodionov)
        statement = build_payroll_statement(db_session, [rodionov], [], YEAR, MONTH)
        row = statement.rows[0]
        assert row.distribution_source == "guard_post"
        assert row.accrued_total == Decimal("75230")
        assert row.guard_tax_amount == Decimal("10092.00")
        amounts = {d.company_id: d.amount for d in row.distribution}
        assert amounts == {
            companies["ZMO"].id: Decimal("29862.70"),
            companies["EKS"].id: Decimal("51193.20"),
            companies["SEC"].id: Decimal("4266.10"),
        }
        assert row.distribution_total == Decimal("85322.00")
        # Разница с начисленным — налог, а не «нераспределённый остаток».
        assert row.unallocated_remainder == _ZERO

    def test_guard_row_without_official_payout_unchanged(
        self, db_session, gbr_place, rodionov
    ):
        _karaulov(db_session, gbr_place, rodionov, official_h1="0")
        row = build_payroll_statement(db_session, [rodionov], [], YEAR, MONTH).rows[0]
        assert row.guard_tax_amount == _ZERO
        assert row.distribution_total == row.accrued_total == Decimal("75230")


class TestExcel:
    def test_tax_column_explains_the_difference(
        self, db_session, gbr_place, rodionov
    ):
        _karaulov(db_session, gbr_place, rodionov)
        month = build_guard_month(db_session, _admin(db_session), YEAR, MONTH)
        ws = load_workbook(BytesIO(generate_guard_timesheet_excel(db_session, month))).active
        header = [c.value for c in ws[4]]
        assert header[-2] == "Итого разбивка"
        assert header[-1] == "Налоги (входят в разбивку)"
        row = next(
            [c.value for c in r] for r in ws.iter_rows()
            if r[2].value == "Караулов Олег Петрович"
        )
        assert row[-2] == 85322.0
        assert row[-1] == 10092.0
        assert 29862.7 in row and 51193.2 in row and 4266.1 in row


# ── Прочие подразделения: ни на рубль ────────────────────────────────────────

AUG_WEEKENDS = "1,2,8,9,15,16,22,23,29,30"
AUG_WORKDAYS = [d for d in range(1, 32) if d not in {int(x) for x in AUG_WEEKENDS.split(",")}]


class TestOtherDepartmentsUntouched:
    """Налог — только вахта. Обычная строка рядом со строкой вахты с налогом
    получает те же суммы при любой ставке, и они равны расчёту по «итого
    начислено» (floor до тысячи)."""

    @pytest.fixture
    def ordinary(self, db_session, other_dept, companies) -> Employee:
        schedule = Schedule(name="5/2", hours_per_shift=8, schedule_type="weekday",
                            is_active=True)
        db_session.add(schedule)
        db_session.add(ProductionCalendar(
            year=YEAR, source="manual",
            data={"year": YEAR, "months": [{"month": MONTH, "days": AUG_WEEKENDS}]},
        ))
        db_session.commit()
        emp = Employee(full_name="Обычный Инженер", tab_number="T-0900", is_active=True,
                       rate=Decimal("152381"), schedule_id=schedule.id,
                       default_company_id=companies["ZMO"].id,
                       department_id=other_dept.id)
        db_session.add(emp)
        db_session.commit()
        position = emp.primary_position
        for day in AUG_WORKDAYS:
            db_session.add(TimesheetEntry(
                employee_id=emp.id, position_id=position.id,
                work_date=date(YEAR, MONTH, day), company_id=companies["ZMO"].id, hours=8,
            ))
        for code, percent in (("ZMO", "35"), ("EKS", "60"), ("SEC", "5")):
            db_session.add(EmployeeCompanyShare(
                employee_id=emp.id, position_id=position.id,
                company_id=companies[code].id, percent=Decimal(percent),
            ))
        db_session.commit()
        return emp

    def _ordinary_row(self, db, ordinary, guard_person):
        entries = db.query(TimesheetEntry).filter(
            TimesheetEntry.employee_id == ordinary.id
        ).all()
        statement = build_payroll_statement(
            db, [ordinary, guard_person], entries, YEAR, MONTH
        )
        return next(r for r in statement.rows if r.employee_id == ordinary.id), statement

    def test_distribution_is_the_same_at_any_tax_rate(
        self, db_session, ordinary, gbr_place, rodionov, companies
    ):
        _karaulov(db_session, gbr_place, rodionov)
        results = []
        for rate in ("0", "40", "100"):
            set_employer_tax_percent(db_session, Decimal(rate))
            db_session.commit()
            row, _ = self._ordinary_row(db_session, ordinary, rodionov)
            assert row.accrued_total == Decimal("152381")  # не пустой расчёт
            results.append((
                row.accrued_total, row.distribution_total, row.unallocated_remainder,
                {d.company_id: d.amount for d in row.distribution},
            ))
        assert results[0] == results[1] == results[2]

    def test_base_is_accrued_and_floor_to_thousands(
        self, db_session, ordinary, gbr_place, rodionov, companies
    ):
        _karaulov(db_session, gbr_place, rodionov)
        row, statement = self._ordinary_row(db_session, ordinary, rodionov)
        assert row.guard_tax_amount == _ZERO
        assert row.accrued_total == Decimal("152381")
        # 152 381 по 35/60/5 → floor + раздача тысяч: как до налогов.
        assert {d.company_id: d.amount for d in row.distribution} == {
            companies["ZMO"].id: Decimal("53000"),
            companies["EKS"].id: Decimal("91000"),
            companies["SEC"].id: Decimal("8000"),
        }
        assert row.distribution_total == Decimal("152000")
        assert row.unallocated_remainder == Decimal("381")
        # Итог ведомости: остаток — только от обычной строки, налог в него не лезет.
        assert statement.total_unallocated_remainder == Decimal("381")


# ── Округление «к выплате» вверх до 500 ₽ по половинам ────────────────────────

class TestPayoutRounding:
    """Правка заказчика: «к выплате» вахты вверх до 500 ₽, каждая половина
    отдельно; ноль и минус не округляются; разнесение не меняется."""

    @pytest.mark.parametrize("exact, rounded", [
        ("49230", "49500"), ("49500", "49500"), ("49500.01", "50000"),
        ("0.01", "500"), ("0", "0"), ("-350", "-350"),
    ])
    def test_round_up_to_500(self, exact, rounded):
        from app.services.guard_payroll import round_guard_payout
        assert round_guard_payout(Decimal(exact)) == Decimal(rounded)

    def test_each_half_is_rounded_separately(self):
        row = calculate_guard_row(
            pay_type="per_shift", rate=Decimal("5000"), year=YEAR, month=MONTH,
            days=_days(1, 31), premium={1: Decimal("230"), 2: Decimal("230")},
            official={1: Decimal("25000"), 2: Decimal("25000")}, tax_percent=FORTY,
        )
        # 1-я: 75 230 − 25 000 = 50 230 → 50 500; 2-я: 80 230 − 25 000 = 55 230 → 55 500.
        assert row.halves[1].net_payout == Decimal("50500")
        assert row.halves[2].net_payout == Decimal("55500")
        # Месяц — сумма округлённых половин, а не округление суммы (105 460 → 105 500).
        assert row.net_payout == Decimal("106000")
        assert row.net_payout_exact == Decimal("105460")
        assert row.rounding_tail == Decimal("-540")

    def test_accrued_tax_and_distribution_are_not_rounded(self, db_session, gbr_place,
                                                           rodionov, companies):
        _karaulov(db_session, gbr_place, rodionov, official_h1="25000")
        row = build_guard_month(db_session, _admin(db_session), YEAR, MONTH) \
            .zones[0].cards[0].rows[0]
        assert row.net_payout_exact == Decimal("50230")
        assert row.net_payout == Decimal("50500")
        assert row.accrued == Decimal("75230")
        assert row.tax == Decimal("10000.00")
        assert row.distribution_base == Decimal("85230.00")
        assert sum(row.distribution.values()) == Decimal("85230.00")

    def test_statement_row_carries_exact_and_tail(self, db_session, gbr_place, rodionov):
        _karaulov(db_session, gbr_place, rodionov, official_h1="25000")
        row = build_payroll_statement(db_session, [rodionov], [], YEAR, MONTH).rows[0]
        assert row.net_payout == Decimal("50500")
        assert row.net_payout_exact == Decimal("50230")
        assert row.rounding_tail == Decimal("-270")
        assert row.distribution_total == Decimal("85230.00")
