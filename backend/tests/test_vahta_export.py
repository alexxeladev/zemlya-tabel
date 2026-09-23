"""Excel-выгрузка табеля вахты (task_vahta п.13).

Формат снят с образца `Табель зп блок охрана август.xlsx`; единственное
согласованное отличие — ОДИН лист на месяц вместо двух листов по половинам,
поэтому суммы половин проверяются отдельной строкой внизу.
"""
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy.orm import Session

from app.models.position_terms import TERMS_BEGINNING
from app.models.companies import Company
from app.models.departments import Department
from app.models.employees import Employee
from app.models.guard_posts import (
    GuardCrew,
    GuardCrewShare,
    GuardPost,
    GuardSite,
    GuardSiteShare,
    GuardZone,
)
from app.services.position_terms import set_effective_from
from app.services.guard_duty import create_assignment
from app.services.guard_export import generate_guard_timesheet_excel
from app.services.guard_month import build_guard_month

YEAR, MONTH = 2026, 8
FIRST_HALF = set(range(1, 16))


@pytest.fixture
def setup(db_session: Session):
    companies = {}
    for i, (code, name) in enumerate(
        [("ZMO", "ЗМО"), ("EKS", "Эксплуатация"), ("SEC", "Секьюрити")], start=1
    ):
        c = Company(code=code, name=name, is_active=True, sort_order=i)
        db_session.add(c)
        companies[code] = c
    dept = Department(
        name="СБ Охрана", code="SEC-OHR", is_active=True, is_guard_department=True
    )
    db_session.add(dept)
    db_session.commit()

    # Зона обслуживания — верхний уровень группировки.
    zone = GuardZone(name="Зона 1", department_id=dept.id)
    db_session.add(zone)
    db_session.commit()

    # Выездной экипаж ГБР зоны со своим мультикомпанийным раскладом.
    crew = GuardCrew(name="1 экипаж", zone_id=zone.id, shift_rate=Decimal("5000"))
    db_session.add(crew)
    db_session.commit()
    for code, percent in [("ZMO", "35"), ("EKS", "60"), ("SEC", "5")]:
        db_session.add(GuardCrewShare(
            crew_id=crew.id, company_id=companies[code].id, percent=Decimal(percent),
        ))

    # Охраняемый объект с постом — чтобы в выгрузке были карточки обоих видов.
    site = GuardSite(name="Green Wood", zone_id=zone.id, shift_rate=Decimal("4000"))
    db_session.add(site)
    db_session.commit()
    db_session.add(GuardSiteShare(
        site_id=site.id, company_id=companies["EKS"].id, percent=Decimal("100"),
    ))
    post = GuardPost(site_id=site.id, name="GW 1")
    db_session.add(post)
    db_session.commit()

    actor = Employee(
        full_name="Админ", email="a@example.com", role="admin", is_active=True
    )
    emp = Employee(
        full_name="Караулов Олег Петрович", tab_number="0000-90274", is_active=True
    )
    db_session.add_all([actor, emp])
    db_session.commit()
    position = emp.ensure_primary_position()
    position.department_id = dept.id
    db_session.commit()

    assignment = create_assignment(
        db_session, year=YEAR, month=MONTH, place=crew,
        position=position, days=FIRST_HALF,
    )
    assignment.premium_h1 = Decimal("230")
    # «Официальный» — свойство РАБОЧЕГО МЕСТА (task_guard_form_rate_official).
    # Зарплата здесь не задаётся намеренно: проверяются колонки выгрузки, а не
    # налог, и цифры остаются теми же, что в образце заказчика.
    set_effective_from(position, TERMS_BEGINNING)  # версии условий: «с начала»
    position.is_official = True
    db_session.commit()
    return {
        "actor": actor, "companies": companies, "employee": emp,
        "zone": zone, "crew": crew, "site": site, "post": post,
    }


def _sheet(db_session, actor):
    month = build_guard_month(db_session, actor, YEAR, MONTH)
    content = generate_guard_timesheet_excel(db_session, month)
    wb = openpyxl.load_workbook(BytesIO(content))
    return wb.active


def _row_values(ws, needle: str) -> list:
    for row in ws.iter_rows(values_only=True):
        if any(str(c) == needle for c in row if c is not None):
            return list(row)
    raise AssertionError(f"строка {needle!r} не найдена")


class TestExport:
    def test_file_opens_and_has_one_sheet(self, db_session, setup):
        month = build_guard_month(db_session, setup["actor"], YEAR, MONTH)
        wb = openpyxl.load_workbook(BytesIO(generate_guard_timesheet_excel(db_session, month)))
        assert len(wb.worksheets) == 1

    def test_headers_follow_the_sample(self, db_session, setup):
        ws = _sheet(db_session, setup["actor"])
        headers = [c.value for c in ws[4]]
        for title in [
            "№", "ФИО", "Должность", "Смена/руб", "КП", "Кол-во смен", "Зарплата",
            "Трудоустройство", "Премия", "Штраф", "Оф. выплата", "К выплате",
            "Примечания", "Итого разбивка",
        ]:
            assert title in headers, title

    def test_day_grid_covers_the_month(self, db_session, setup):
        ws = _sheet(db_session, setup["actor"])
        headers = [c.value for c in ws[4]]
        assert all(day in headers for day in range(1, 32))

    def test_cards_are_labelled_by_kind(self, db_session, setup):
        """Зона, экипаж ГБР и объект — разные вещи, и в выгрузке подписаны."""
        ws = _sheet(db_session, setup["actor"])
        texts = [
            str(v)
            for row in ws.iter_rows(values_only=True)
            for v in row
            if v is not None
        ]
        assert any("ЗОНА · Зона 1" in t for t in texts)
        assert any("ЭКИПАЖ ГБР · 1 экипаж" in t for t in texts)
        assert any("ОБЪЕКТ · Green Wood" in t for t in texts)

    def test_company_columns_present(self, db_session, setup):
        ws = _sheet(db_session, setup["actor"])
        headers = [str(c.value) for c in ws[4]]
        for name in ("ЗМО", "Эксплуатация", "Секьюрити"):
            assert name in headers

    def test_row_numbers_match_the_calculation(self, db_session, setup):
        ws = _sheet(db_session, setup["actor"])
        row = _row_values(ws, "Караулов Олег Петрович")
        assert 75000 in row      # зарплата
        assert 230 in row        # премия
        assert 15 in row         # кол-во смен
        assert 26330.5 in row    # ЗМО 35 %
        assert 45138 in row      # Эксплуатация 60 %
        assert 3761.5 in row     # Секьюрити 5 %
        assert 75230 in row      # итого разбивка
        assert "Официальный" in row

    def test_empty_slot_does_not_crash(self, db_session, setup):
        create_assignment(
            db_session, year=YEAR, month=MONTH, place=setup["post"], position=None,
            days=set(),
        )
        db_session.commit()
        ws = _sheet(db_session, setup["actor"])
        assert _row_values(ws, "— вакансия —")

    def test_half_totals_are_printed(self, db_session, setup):
        ws = _sheet(db_session, setup["actor"])
        texts = [
            str(value)
            for row in ws.iter_rows(values_only=True)
            for value in row
            if value is not None
        ]
        assert any("Половина 1 (1–15)" in t for t in texts)
        assert any("Половина 2 (16–31)" in t for t in texts)

    def test_totals_row(self, db_session, setup):
        ws = _sheet(db_session, setup["actor"])
        row = _row_values(ws, "ИТОГО")
        assert 75230 in row
