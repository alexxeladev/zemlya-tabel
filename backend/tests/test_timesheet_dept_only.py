"""
Табель открывается только по ОДНОМУ отделу (task_timesheet_dept_only).

Режим «все отделы» снят с экрана, и группе «Без отдела» понадобилось своё
значение фильтра: раньше она приходила только заодно со всеми отделами, а без
неё сотрудники без отдела и их периоды стали бы недостижимы с экрана.

Проверяется КОНТРАКТ бэка, на который опирается экран:
  · `?department_id=none` отдаёт РОВНО группу «Без отдела» — и людей, и период;
  · чужой отдел и негодное значение не превращаются молча в «показать всех»;
  · менеджеру и табельщику группа недоступна (403), их выдача не меняется;
  · то же значение работает в расчёте, премиях, Т-13 и автозаполнении.
"""
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models.companies import Company
from app.models.departments import Department
from app.models.employees import Employee
from app.models.schedules import Schedule
from app.services.positions import (
    NO_DEPARTMENT,
    normalize_department_filter,
)
from tests.conftest import get_token

YEAR, MONTH = 2026, 7


# ── Фикстуры ──────────────────────────────────────────────────────────────────

@pytest.fixture
def company(db_session: Session) -> Company:
    c = Company(name="ООО «Земля МО»", code="ZMO", is_active=True)
    db_session.add(c)
    db_session.commit()
    db_session.refresh(c)
    return c


@pytest.fixture
def schedule(db_session: Session) -> Schedule:
    s = Schedule(name="5/2", schedule_type="weekday", hours_per_shift=8, is_active=True)
    db_session.add(s)
    db_session.commit()
    db_session.refresh(s)
    return s


@pytest.fixture
def dept(db_session: Session, company: Company) -> Department:
    d = Department(name="ИТО", code="ITO", head_company_id=company.id, is_active=True)
    db_session.add(d)
    db_session.commit()
    db_session.refresh(d)
    return d


def _worker(db_session, name, tab, dept_id, company, schedule) -> Employee:
    emp = Employee(
        full_name=name,
        tab_number=tab,
        position="Инженер",
        department_id=dept_id,
        default_company_id=company.id,
        schedule_id=schedule.id,
        rate=Decimal("50000"),
        is_active=True,
    )
    db_session.add(emp)
    db_session.commit()
    db_session.refresh(emp)
    return emp


@pytest.fixture
def in_dept(db_session, dept, company, schedule) -> Employee:
    return _worker(db_session, "Иванов Иван", "T-001", dept.id, company, schedule)


@pytest.fixture
def no_dept(db_session, company, schedule) -> Employee:
    return _worker(db_session, "Беспалов Борис", "T-002", None, company, schedule)


@pytest.fixture
def admin_token(client: TestClient, admin_user: Employee) -> str:
    return get_token(client, "admin@example.com", "admin123")


@pytest.fixture
def manager_token(client: TestClient, db_session, manager_user: Employee, dept) -> str:
    manager_user.managed_departments = [dept]
    db_session.commit()
    return get_token(client, "manager@example.com", "manager123")


@pytest.fixture
def timekeeper_token(client: TestClient, db_session, dept) -> str:
    emp = Employee(
        full_name="Табельщик", email="tk@example.com",
        hashed_password=hash_password("tk123456"), role="timekeeper",
        is_active=True, must_change_password=False,
    )
    emp.managed_departments = [dept]
    db_session.add(emp)
    db_session.commit()
    return get_token(client, "tk@example.com", "tk123456")


def _hdr(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _month(client, token, **params):
    return client.get(
        f"/api/timesheet/{YEAR}/{MONTH}", headers=_hdr(token), params=params
    )


# ── Разбор значения фильтра (единственное место) ──────────────────────────────

class TestNormalizeDepartmentFilter:
    def test_none_word_is_the_group(self):
        assert normalize_department_filter("none") is NO_DEPARTMENT

    def test_only_lowercase(self):
        """В теле автозаполнения то же значение описано `Literal["none"]`, и
        pydantic отвергает «NONE» до нас. Приняв здесь любой регистр, получили
        бы два разных контракта на одно значение."""
        with pytest.raises(ValueError):
            normalize_department_filter("NONE")

    def test_bool_is_not_a_department(self):
        """`bool` — подкласс `int`: без явного отказа `True` стал бы отделом №1."""
        for bad in (True, False):
            with pytest.raises(ValueError):
                normalize_department_filter(bad)

    def test_number_and_empty(self):
        assert normalize_department_filter("42") == 42
        assert normalize_department_filter(42) == 42
        assert normalize_department_filter(None) is None
        assert normalize_department_filter("") is None

    def test_garbage_is_an_error_not_show_everyone(self):
        """Опечатка не должна молча означать «фильтра нет»."""
        with pytest.raises(ValueError):
            normalize_department_filter("all")


# ── Группа «Без отдела» ───────────────────────────────────────────────────────

class TestNoDepartmentGroup:
    def test_returns_only_employees_without_a_department(
        self, client, admin_token, in_dept, no_dept
    ):
        resp = _month(client, admin_token, department_id="none")
        assert resp.status_code == 200
        ids = {e["id"] for e in resp.json()["employees"]}
        assert ids == {no_dept.id}

    def test_department_id_still_selects_that_department(
        self, client, admin_token, dept, in_dept, no_dept
    ):
        resp = _month(client, admin_token, department_id=dept.id)
        assert resp.status_code == 200
        assert {e["id"] for e in resp.json()["employees"]} == {in_dept.id}

    def test_group_carries_its_period_so_it_can_be_closed(
        self, client, admin_token, in_dept, no_dept
    ):
        """Период группы — с `department_id: null`. Без него «Задачи» вели бы
        в табель, где закрывать нечего."""
        periods = _month(client, admin_token, department_id="none").json()["periods"]
        assert [p["department_id"] for p in periods] == [None]
        assert periods[0]["id"] > 0

    def test_positions_of_the_group_only(self, client, admin_token, in_dept, no_dept):
        by_emp = _month(client, admin_token, department_id="none").json()[
            "positions_by_employee"
        ]
        assert set(by_emp) == {str(no_dept.id)}
        assert all(p["department_id"] is None for p in by_emp[str(no_dept.id)])

    def test_payroll_accepts_the_group(self, client, admin_token, in_dept, no_dept):
        resp = client.get(
            f"/api/timesheet/{YEAR}/{MONTH}/payroll",
            headers=_hdr(admin_token), params={"department_id": "none"},
        )
        assert resp.status_code == 200
        assert {r["employee_id"] for r in resp.json()["employees"]} == {no_dept.id}

    def test_adjustments_accept_the_group(self, client, admin_token, no_dept):
        resp = client.get(
            f"/api/timesheet/{YEAR}/{MONTH}/adjustments",
            headers=_hdr(admin_token), params={"department_id": "none"},
        )
        assert resp.status_code == 200

    def test_quantities_of_the_group_are_empty(self, client, admin_token, no_dept):
        """Количественный показатель — настройка ОТДЕЛА, у группы его быть не может."""
        resp = client.get(
            f"/api/timesheet/{YEAR}/{MONTH}/quantities",
            headers=_hdr(admin_token), params={"department_id": "none"},
        )
        assert resp.status_code == 200
        assert resp.json() == []

    def test_t13_of_one_department_is_titled_by_it_without_a_filter(
        self, client, db_session, manager_token, dept, in_dept
    ):
        """Роль без выбора отдела шлёт пустой фильтр (её набор отделов знает
        сервер), но в файле один отдел — подпись «Все отделы» врала бы. Номер
        отдела из фронта брать нельзя: профиль там кэширован."""
        from io import BytesIO

        from openpyxl import load_workbook

        resp = client.get(
            f"/api/timesheet/{YEAR}/{MONTH}/export/excel", headers=_hdr(manager_token),
        )
        assert resp.status_code == 200
        ws = load_workbook(BytesIO(resp.content)).active
        text = "\n".join(
            str(c.value) for row in ws.iter_rows() for c in row if c.value
        )
        assert dept.name in text
        assert "Все отделы" not in text

    def test_t13_export_is_titled_by_the_group(self, client, admin_token, no_dept):
        resp = client.get(
            f"/api/timesheet/{YEAR}/{MONTH}/export/excel",
            headers=_hdr(admin_token), params={"department_id": "none"},
        )
        assert resp.status_code == 200
        assert resp.content[:2] == b"PK"

    def test_autofill_preview_touches_only_the_group(
        self, client, admin_token, db_session, in_dept, no_dept
    ):
        from app.models.production_calendars import ProductionCalendar

        db_session.add(ProductionCalendar(
            year=YEAR,
            data={"year": YEAR, "months": [
                {"month": MONTH, "days": "4,5,11,12,18,19,25,26"}
            ]},
            source="manual",
        ))
        db_session.commit()
        resp = client.post(
            "/api/timesheet/autofill/preview", headers=_hdr(admin_token),
            json={"year": YEAR, "month": MONTH, "department_id": "none"},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        touched = {c["employee_id"] for c in body["entries_to_create"]}
        assert touched == {no_dept.id}, "автозаполнение вышло за группу"
        assert body["employees_processed"] == 1


# ── Права: группа — только admin/accountant ───────────────────────────────────

class TestGroupAccess:
    def test_manager_gets_403(self, client, manager_token, no_dept):
        assert _month(client, manager_token, department_id="none").status_code == 403

    def test_timekeeper_gets_403(self, client, timekeeper_token, no_dept):
        assert _month(client, timekeeper_token, department_id="none").status_code == 403

    def test_employee_gets_403_not_an_empty_timesheet(
        self, client, db_session, no_dept, schedule, company
    ):
        """Сотруднику группа тоже ни к чему — он видит себя. Без явного отказа
        `visible_positions` отсекала бы у него все места с отделом, и он получал
        бы ПУСТОЙ табель вместо внятного «нельзя»."""
        emp = _worker(db_session, "Сам Себе", "T-009", None, company, schedule)
        emp.email, emp.role = "self@example.com", "employee"
        emp.hashed_password = hash_password("selfpass1")
        emp.must_change_password = False
        db_session.commit()
        token = get_token(client, "self@example.com", "selfpass1")
        assert _month(client, token, department_id="none").status_code == 403
        # без фильтра он по-прежнему видит себя
        assert _month(client, token).status_code == 200

    def test_manager_without_filter_sees_his_department_only(
        self, client, manager_token, in_dept, no_dept
    ):
        """Выдача ролей по отделам не поехала: группа в неё не подмешалась."""
        resp = _month(client, manager_token)
        assert resp.status_code == 200
        assert {e["id"] for e in resp.json()["employees"]} == {in_dept.id}

    def test_manager_payroll_of_the_group_is_403(self, client, manager_token, no_dept):
        resp = client.get(
            f"/api/timesheet/{YEAR}/{MONTH}/payroll",
            headers=_hdr(manager_token), params={"department_id": "none"},
        )
        assert resp.status_code == 403


# ── Негодное значение не превращается в «показать всех» ───────────────────────

class TestBadValue:
    @pytest.mark.parametrize("bad", ["all", "abc", "1,2"])
    def test_rejected_with_422(self, client, admin_token, bad, in_dept, no_dept):
        resp = _month(client, admin_token, department_id=bad)
        assert resp.status_code == 422, bad

    def test_autofill_body_rejects_bool(self, client, admin_token):
        """pydantic приводил `true` к 1 ДО нашего разбора — автозаполнение
        уходило в отдел №1 вместо отказа."""
        resp = client.post(
            "/api/timesheet/autofill/apply", headers=_hdr(admin_token),
            json={"year": YEAR, "month": MONTH, "department_id": True},
        )
        assert resp.status_code == 422

    def test_autofill_body_rejects_uppercase_none(self, client, admin_token):
        """Разбор в теле и в query обязан совпадать — иначе клиент, работающий
        с GET, получает 422 на автозаполнении с тем же значением."""
        resp = client.post(
            "/api/timesheet/autofill/apply", headers=_hdr(admin_token),
            json={"year": YEAR, "month": MONTH, "department_id": "NONE"},
        )
        assert resp.status_code == 422
        assert _month(client, admin_token, department_id="NONE").status_code == 422

    def test_missing_filter_still_means_everyone(
        self, client, admin_token, in_dept, no_dept
    ):
        """Контракт «фильтра нет = все» не менялся: его используют ведомость и
        роли с одним отделом."""
        resp = _month(client, admin_token)
        assert {e["id"] for e in resp.json()["employees"]} == {in_dept.id, no_dept.id}
