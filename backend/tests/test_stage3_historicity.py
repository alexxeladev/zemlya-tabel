"""
Этап 3: историчность расчёта (task_stage3_historicity).

По пунктам приёмки:
  1. миграция не сдвигает сумм — позиция с одной версией «с начала» считается
     так же, как без версий (сверка всей дев-базы — отдельным отчётом);
  2. правка ставки после закрытия не меняет закрытую ведомость и её Excel;
  3. повторное построение закрытой ведомости — те же суммы до копейки;
  4. повышение оклада с середины месяца — с даты повышения (пример на числах);
  5. переоткрытие аннулирует снимок, повторное закрытие создаёт новый;
  6. премии, аванс, займ, проценты, показатель в закрытом периоде и на
     проверке — отказ на бэке во всех точках входа;
  8. история условий видна в карточке позиции.
Плюс: стык версий, совместитель, займ по снимку, дашборд = снимок, вахта.
"""
from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy.orm import Session

from app.models.companies import Company
from app.models.departments import Department
from app.models.employee_adjustments import EmployeeAdjustment
from app.models.employees import Employee
from app.models.loan_deductions import LoanDeduction
from app.models.period_snapshots import PeriodSnapshot
from app.models.position_terms import TERMS_BEGINNING, PositionTerms
from app.models.positions import EmployeePosition
from app.models.production_calendars import ProductionCalendar
from app.models.schedules import Schedule
from app.models.timesheet_entries import TimesheetEntry
from app.models.timesheet_periods import TimesheetPeriod
from app.services.payroll import calculate_position_payroll
from app.services.payroll_statement import build_payroll_statement, build_payroll_summary
from app.services.position_terms import (
    ClosedPeriodError,
    default_effective_from,
    month_segments,
    set_effective_from,
    terms_on,
)
from app.services.timesheet_periods import close_period, reopen_period
from tests.conftest import get_token

# Май 2026: выходные и праздники 1–3, 9–11 не рабочие по реальному календарю;
# здесь упрощённо — 20 рабочих дней по 8 ч, норма 160 ч.
MAY_REAL = {"year": 2026, "months": [{"month": 5, "days": "1,2,3,9,10,16,17,23,24,30,31"}]}
MAY_WORKDAYS = [d for d in range(1, 32) if d not in (1, 2, 3, 9, 10, 16, 17, 23, 24, 30, 31)]


@pytest.fixture
def setup(db_session: Session, admin_user):
    cal = ProductionCalendar(year=2026, data=MAY_REAL, source="manual")
    sched = Schedule(name="5/2", hours_per_shift=8, schedule_type="weekday",
                     work_weekdays=[0, 1, 2, 3, 4])
    zmo = Company(code="ZMO", name="Земля МО")
    kft = Company(code="KFT", name="Комфорт")
    db_session.add_all([cal, sched, zmo, kft])
    db_session.commit()
    dept = Department(name="ИТО", code="ITO", head_company_id=zmo.id)
    db_session.add(dept)
    db_session.commit()
    emp = Employee(
        full_name="Окладник Иван", tab_number="T-0001", rate=Decimal("60000"),
        schedule_id=sched.id, department_id=dept.id, default_company_id=zmo.id,
    )
    db_session.add(emp)
    db_session.commit()
    db_session.refresh(emp)
    for day in MAY_WORKDAYS:
        db_session.add(TimesheetEntry(
            employee_id=emp.id, position_id=emp.primary_position.id,
            company_id=zmo.id, work_date=date(2026, 5, day), hours=8,
        ))
    db_session.commit()
    return {"emp": emp, "dept": dept, "sched": sched, "zmo": zmo, "kft": kft,
            "admin": admin_user}


def _auth(client, admin_user):
    token = get_token(client, "admin@example.com", "admin123")
    return {"Authorization": f"Bearer {token}"}


def _statement(db, emp, year=2026, month=5):
    entries = db.query(TimesheetEntry).all()
    return build_payroll_statement(db, [emp], entries, year, month)


def _period(db, dept, status="pending_review", month=5):
    p = TimesheetPeriod(department_id=dept.id, year=2026, month=month, status=status)
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


def _raise_rate(db, emp, amount, day):
    pos = emp.primary_position
    set_effective_from(pos, day)
    pos.rate = Decimal(amount)
    db.commit()
    db.refresh(emp)


# ── Часть 1: версии условий ───────────────────────────────────────────────────

class TestTermsVersions:
    def test_new_position_gets_version_from_beginning(self, db_session, setup):
        pos = setup["emp"].primary_position
        assert len(pos.terms_versions) == 1
        v = pos.terms_versions[0]
        assert v.effective_from == TERMS_BEGINNING
        assert v.rate == Decimal("60000.00")
        assert v.schedule_id == setup["sched"].id

    def test_one_version_is_identical_to_raw_position(self, db_session, setup):
        """Миграция: одна версия «с начала» не меняет ни одной суммы."""
        emp = setup["emp"]
        pos = emp.primary_position
        entries = db_session.query(TimesheetEntry).all()
        versioned = calculate_position_payroll(emp, pos, entries, MAY_REAL, 2026, 5)

        class Raw:  # та же позиция без истории — как до этапа 3
            def __init__(self, p):
                self._p = p

            def __getattr__(self, name):
                if name == "terms_versions":
                    return []
                return getattr(self._p, name)

        raw = calculate_position_payroll(emp, Raw(pos), entries, MAY_REAL, 2026, 5)
        assert versioned == raw

    def test_change_without_date_starts_next_month(self, db_session, setup):
        pos = setup["emp"].primary_position
        pos.rate = Decimal("70000")
        db_session.commit()
        assert [v.effective_from for v in pos.terms_versions] == [
            TERMS_BEGINNING, default_effective_from(),
        ]
        # Зеркало — последняя версия.
        assert pos.rate == Decimal("70000.00")

    def test_mid_month_raise_is_paid_from_its_date(self, db_session, setup):
        """Пример на числах: оклад 60 000 → 90 000 с 15 мая, отработан весь май.

        1–14 мая — 9 рабочих дней × 8 ч = 72 ч по старому окладу,
        15–31 мая — 11 дней × 8 ч = 88 ч по новому, норма 160 ч:
        60 000 × 72/160 + 90 000 × 88/160 = 27 000 + 49 500 = 76 500 ₽.
        Раньше (оклад на весь месяц) было бы 90 000 ₽.
        """
        emp = setup["emp"]
        _raise_rate(db_session, emp, "90000", date(2026, 5, 15))
        entries = db_session.query(TimesheetEntry).all()
        p = calculate_position_payroll(emp, emp.primary_position, entries, MAY_REAL, 2026, 5)
        assert p.base_amount == Decimal("76500")
        assert p.total_amount == Decimal("76500")
        assert p.total_hours == Decimal("160")
        assert p.rate == Decimal("90000")  # ставка на конец месяца

    def test_boundary_day_belongs_to_new_version(self, db_session, setup):
        """Стык: 14-е — последний день старой версии, 15-е — первый новой."""
        emp = setup["emp"]
        _raise_rate(db_session, emp, "90000", date(2026, 5, 15))
        pos = emp.primary_position
        assert terms_on(pos, date(2026, 5, 14)).rate == Decimal("60000.00")
        assert terms_on(pos, date(2026, 5, 15)).rate == Decimal("90000.00")
        segs = month_segments(pos, 2026, 5)
        assert [(s, e) for s, e, _ in segs] == [
            (date(2026, 5, 1), date(2026, 5, 14)),
            (date(2026, 5, 15), date(2026, 5, 31)),
        ]

    def test_raise_from_first_of_month_is_whole_month(self, db_session, setup):
        emp = setup["emp"]
        _raise_rate(db_session, emp, "90000", date(2026, 5, 1))
        entries = db_session.query(TimesheetEntry).all()
        p = calculate_position_payroll(emp, emp.primary_position, entries, MAY_REAL, 2026, 5)
        assert p.base_amount == Decimal("90000")
        assert len(month_segments(emp.primary_position, 2026, 5)) == 1

    def test_earlier_change_propagates_to_later_versions(self, db_session, setup):
        """Поле, изменённое с даты D, действует с D и дальше — более поздняя
        версия его не откатывает."""
        emp = setup["emp"]
        pos = emp.primary_position
        set_effective_from(pos, date(2026, 7, 1))
        pos.overtime_coefficient = Decimal("1")
        db_session.commit()
        set_effective_from(pos, date(2026, 6, 1))
        pos.rate = Decimal("80000")
        db_session.commit()
        assert terms_on(pos, date(2026, 6, 10)).rate == Decimal("80000.00")
        assert terms_on(pos, date(2026, 7, 10)).rate == Decimal("80000.00")
        assert terms_on(pos, date(2026, 7, 10)).overtime_coefficient == Decimal("1.00")
        assert terms_on(pos, date(2026, 6, 10)).overtime_coefficient == Decimal("1.50")
        assert pos.rate == Decimal("80000.00")

    def test_earlier_change_keeps_planned_later_change(self, db_session, setup):
        """Запланировано 100 000 с 1 ноября; потом вводят 90 000 с 15 октября.
        Ноябрьское повышение не затирается: 15.10–31.10 — 90 000, с 1.11 — 100 000."""
        pos = setup["emp"].primary_position
        set_effective_from(pos, date(2026, 11, 1))
        pos.rate = Decimal("100000")
        db_session.commit()
        set_effective_from(pos, date(2026, 10, 15))
        pos.rate = Decimal("90000")
        db_session.commit()
        assert terms_on(pos, date(2026, 10, 14)).rate == Decimal("60000.00")
        assert terms_on(pos, date(2026, 10, 15)).rate == Decimal("90000.00")
        assert terms_on(pos, date(2026, 11, 1)).rate == Decimal("100000.00")
        assert pos.rate == Decimal("100000.00")  # зеркало — последняя версия

    def test_moonlighter_positions_have_own_history(self, db_session, setup):
        emp = setup["emp"]
        emp.positions.append(EmployeePosition(
            title="Электрик", rate=Decimal("30000"), schedule_id=setup["sched"].id,
            company_id=setup["kft"].id, department_id=setup["dept"].id,
        ))
        db_session.commit()
        second = emp.positions[1]
        set_effective_from(second, date(2026, 5, 15))
        second.rate = Decimal("45000")
        db_session.commit()
        assert len(emp.primary_position.terms_versions) == 1
        assert len(second.terms_versions) == 2

    def test_version_in_closed_month_rejected(self, db_session, setup):
        _period(db_session, setup["dept"], status="closed")
        pos = setup["emp"].primary_position
        set_effective_from(pos, date(2026, 5, 20))
        pos.rate = Decimal("90000")
        with pytest.raises(ClosedPeriodError):
            db_session.commit()
        db_session.rollback()

    def test_version_in_pending_month_rejected(self, db_session, setup):
        _period(db_session, setup["dept"], status="pending_review")
        pos = setup["emp"].primary_position
        set_effective_from(pos, date(2026, 5, 20))
        pos.rate = Decimal("90000")
        with pytest.raises(ClosedPeriodError):
            db_session.commit()
        db_session.rollback()

    def test_rewrite_from_beginning_rejected_when_closed_months_exist(self, db_session, setup):
        _period(db_session, setup["dept"], status="closed")
        pos = setup["emp"].primary_position
        set_effective_from(pos, TERMS_BEGINNING)
        pos.rate = Decimal("90000")
        with pytest.raises(ClosedPeriodError):
            db_session.commit()
        db_session.rollback()

    def test_api_takes_effective_date_and_rejects_closed(self, client, db_session, setup):
        h = _auth(client, setup["admin"])
        emp, pos = setup["emp"], setup["emp"].primary_position
        r = client.patch(f"/api/employees/{emp.id}/positions/{pos.id}", headers=h,
                         json={"rate": "90000", "terms_effective_from": "2026-05-15"})
        assert r.status_code == 200, r.text
        _period(db_session, setup["dept"], status="closed", month=6)
        r = client.patch(f"/api/employees/{emp.id}/positions/{pos.id}", headers=h,
                         json={"rate": "95000", "terms_effective_from": "2026-06-10"})
        assert r.status_code == 409
        assert "закрыт" in r.json()["detail"]

    def test_history_in_position_card(self, client, db_session, setup):
        h = _auth(client, setup["admin"])
        emp, pos = setup["emp"], setup["emp"].primary_position
        client.patch(f"/api/employees/{emp.id}/positions/{pos.id}", headers=h,
                     json={"rate": "90000", "terms_effective_from": "2026-05-15"})
        r = client.get(f"/api/employees/{emp.id}/positions/{pos.id}/terms", headers=h)
        assert r.status_code == 200
        body = r.json()
        assert [v["effective_from"] for v in body["versions"]] == [None, "2026-05-15"]
        assert body["versions"][1]["changed"] == ["Оклад"]
        assert body["versions"][1]["created_by_name"] == "Test Admin"
        assert body["default_effective_from"] == default_effective_from().isoformat()


# ── Часть 2: снимок закрытого периода ─────────────────────────────────────────

class TestSnapshot:
    def test_close_takes_snapshot_and_rate_change_does_not_move_it(self, db_session, setup):
        emp = setup["emp"]
        before = _statement(db_session, emp)
        period = _period(db_session, setup["dept"])
        close_period(db_session, period, setup["admin"])
        assert db_session.query(PeriodSnapshot).count() == 1

        # Правка «с начала» после закрытия запрещена, поэтому ставку меняем
        # так, как её можно было бы поменять до этапа 3 — мимо истории, прямо
        # в версии: снимок обязан этого не заметить.
        db_session.query(PositionTerms).update({"rate": Decimal("120000")})
        db_session.commit()
        db_session.expire_all()
        after = _statement(db_session, emp)
        assert after.rows[0].accrued_total == before.rows[0].accrued_total == Decimal("60000")
        assert after.model_dump() == before.model_dump()

    def test_closed_statement_rebuilds_identically(self, db_session, setup):
        emp = setup["emp"]
        close_period(db_session, _period(db_session, setup["dept"]), setup["admin"])
        a = _statement(db_session, emp).model_dump(mode="json")
        b = _statement(db_session, emp).model_dump(mode="json")
        assert a == b

    def test_excel_of_closed_statement_ignores_share_change(self, client, db_session, setup):
        h = _auth(client, setup["admin"])
        emp = setup["emp"]
        close_period(db_session, _period(db_session, setup["dept"]), setup["admin"])
        before = client.get("/api/timesheet/2026/5/statement", headers=h).json()
        # Дефолт отдела не версионируется — закрытый месяц защищает снимок.
        r = client.put(f"/api/departments/{setup['dept'].id}/company-shares", headers=h,
                       json={"shares": [{"company_id": setup["kft"].id, "percent": "100"}]})
        assert r.status_code == 200, r.text
        after = client.get("/api/timesheet/2026/5/statement", headers=h).json()
        assert after["rows"] == before["rows"]
        xlsx = client.get("/api/timesheet/2026/5/statement/export/excel", headers=h)
        assert xlsx.status_code == 200
        from io import BytesIO

        from openpyxl import load_workbook
        sheet = load_workbook(BytesIO(xlsx.content)).active
        values = [c for row in sheet.iter_rows(values_only=True) for c in row if c is not None]
        assert 60000 in values
        assert emp.id == before["rows"][0]["employee_id"]

    def test_reopen_drops_snapshot_and_reclose_takes_new(self, db_session, setup):
        emp = setup["emp"]
        period = _period(db_session, setup["dept"])
        close_period(db_session, period, setup["admin"])
        first = db_session.query(PeriodSnapshot).one()
        assert Decimal(first.statement_rows[0]["accrued_total"]) == Decimal("60000")

        reopen_period(db_session, period, setup["admin"], "исправить оклад")
        assert db_session.query(PeriodSnapshot).count() == 0
        # Открытый месяц снова считается на лету: оклад с 1 мая меняет суммы.
        _raise_rate(db_session, emp, "90000", date(2026, 5, 1))
        assert _statement(db_session, emp).rows[0].accrued_total == Decimal("90000")

        period.status = "pending_review"
        db_session.commit()
        close_period(db_session, period, setup["admin"])
        # Новый снимок (id в SQLite переиспользуется — сверяем содержимое).
        snap = db_session.query(PeriodSnapshot).one()
        assert Decimal(snap.statement_rows[0]["accrued_total"]) == Decimal("90000")
        used = snap.terms_used[str(emp.primary_position.id)]
        assert used[0]["rate"] == "90000.00" and used[0]["from"] == "2026-05-01"

    def test_snapshot_keeps_distribution(self, db_session, setup):
        emp = setup["emp"]
        close_period(db_session, _period(db_session, setup["dept"]), setup["admin"])
        dist = _statement(db_session, emp).rows[0].distribution
        assert [(d.company_id, d.amount) for d in dist] == [(setup["zmo"].id, Decimal("60000"))]

    def test_dashboard_reads_closed_month_from_snapshot(self, client, db_session, setup):
        h = _auth(client, setup["admin"])
        close_period(db_session, _period(db_session, setup["dept"]), setup["admin"])
        before = client.get("/api/dashboard/2026/5", headers=h).json()["payroll"]["total"]
        db_session.query(PositionTerms).update({"rate": Decimal("120000")})
        db_session.commit()
        # Правка мимо ORM кэш не сбрасывает — сбрасываем как после ручного SQL.
        from app.services.dashboard_cache import drop_cache
        drop_cache(db_session)
        db_session.commit()
        after = client.get("/api/dashboard/2026/5", headers=h).json()["payroll"]["total"]
        assert Decimal(str(after)) == Decimal(str(before)) == Decimal("60000")

    def test_snapshot_read_is_not_a_recompute(self, db_session, setup, monkeypatch):
        """Закрытая ведомость не зовёт расчёт вовсе — только читает снимок."""
        emp = setup["emp"]
        close_period(db_session, _period(db_session, setup["dept"]), setup["admin"])
        import app.services.payroll_statement as ps

        def boom(*a, **k):
            raise AssertionError("расчёт закрытого месяца не должен вызываться")

        monkeypatch.setattr(ps, "calculate_position_payroll", boom)
        assert _statement(db_session, emp).rows[0].accrued_total == Decimal("60000")


# ── Часть 2: запрет правок закрытого периода ──────────────────────────────────

@pytest.fixture(params=["closed", "pending_review"])
def locked(request, db_session, setup):
    setup["dept"].uses_quantity_distribution = True
    db_session.commit()
    _period(db_session, setup["dept"], status=request.param)
    return request.param


class TestClosedPeriodBan:
    def test_premium_create_rejected(self, client, locked, setup):
        h = _auth(client, setup["admin"])
        r = client.post("/api/timesheet/adjustments", headers=h, json={
            "employee_id": setup["emp"].id, "year": 2026, "month": 5,
            "kind": "premium", "amount": "5000", "reason": "за объект",
        })
        assert r.status_code == 409, r.text

    def test_premium_delete_rejected(self, client, db_session, setup):
        adj = EmployeeAdjustment(employee_id=setup["emp"].id, year=2026, month=5,
                                 kind="advance", amount=Decimal("1000"), reason="аванс")
        db_session.add(adj)
        db_session.commit()
        _period(db_session, setup["dept"], status="closed")
        h = _auth(client, setup["admin"])
        r = client.delete(f"/api/timesheet/adjustments/{adj.id}", headers=h)
        assert r.status_code == 409
        assert db_session.get(EmployeeAdjustment, adj.id) is not None

    def test_loan_override_rejected(self, client, db_session, locked, setup):
        emp = setup["emp"]
        emp.loan_amount, emp.loan_term_months, emp.loan_start_date = (
            Decimal("12000"), 12, date(2026, 1, 1))
        db_session.commit()
        h = _auth(client, setup["admin"])
        r = client.post("/api/timesheet/loan-override", headers=h, json={
            "employee_id": emp.id, "year": 2026, "month": 5, "actual_amount": "0",
        })
        assert r.status_code == 409, r.text

    def test_loan_override_delete_rejected(self, client, db_session, setup):
        emp = setup["emp"]
        emp.loan_amount, emp.loan_term_months, emp.loan_start_date = (
            Decimal("12000"), 12, date(2026, 1, 1))
        db_session.add(LoanDeduction(employee_id=emp.id, year=2026, month=5,
                                     planned_amount=Decimal("1000"), actual_amount=Decimal("0")))
        db_session.commit()
        _period(db_session, setup["dept"], status="closed")
        h = _auth(client, setup["admin"])
        r = client.delete(f"/api/timesheet/loan-override/{emp.id}/2026/5", headers=h)
        assert r.status_code == 409

    def test_distribution_override_rejected(self, client, locked, setup):
        h = _auth(client, setup["admin"])
        r = client.put("/api/timesheet/distribution", headers=h, json={
            "employee_id": setup["emp"].id, "year": 2026, "month": 5,
            "shares": [{"company_id": setup["kft"].id, "percent": "100"}],
        })
        assert r.status_code == 409, r.text

    def test_distribution_override_delete_rejected(self, client, db_session, setup):
        from app.models.company_shares import CompanyShareOverride
        db_session.add(CompanyShareOverride(
            employee_id=setup["emp"].id, position_id=setup["emp"].primary_position.id,
            company_id=setup["kft"].id, year=2026, month=5, percent=Decimal("100"),
        ))
        db_session.commit()
        _period(db_session, setup["dept"], status="closed")
        h = _auth(client, setup["admin"])
        r = client.delete(f"/api/timesheet/distribution/{setup['emp'].id}/2026/5", headers=h)
        assert r.status_code == 409
        assert db_session.query(CompanyShareOverride).count() == 1

    def test_quantities_rejected(self, client, locked, setup):
        h = _auth(client, setup["admin"])
        r = client.put("/api/timesheet/quantities", headers=h, json={
            "department_id": setup["dept"].id, "year": 2026, "month": 5,
            "items": [{"company_id": setup["kft"].id, "part1": 3, "part2": 0}],
        })
        assert r.status_code == 409, r.text

    def test_card_shares_from_closed_month_rejected(self, client, locked, setup):
        h = _auth(client, setup["admin"])
        r = client.put(f"/api/employees/{setup['emp'].id}/company-shares", headers=h, json={
            "shares": [{"company_id": setup["kft"].id, "percent": "100"}],
            "effective_from": "2026-05-01",
        })
        assert r.status_code == 409, r.text

    def test_draft_month_still_editable(self, client, db_session, setup):
        _period(db_session, setup["dept"], status="draft")
        h = _auth(client, setup["admin"])
        r = client.post("/api/timesheet/adjustments", headers=h, json={
            "employee_id": setup["emp"].id, "year": 2026, "month": 5,
            "kind": "premium", "amount": "5000", "reason": "за объект",
        })
        assert r.status_code in (200, 201), r.text

    def test_service_layer_is_guarded_too(self, db_session, setup):
        """Запрет — в сессии, а не в роутере: запись мимо API тоже отклоняется."""
        _period(db_session, setup["dept"], status="closed")
        db_session.add(EmployeeAdjustment(employee_id=setup["emp"].id, year=2026, month=5,
                                          kind="premium", amount=Decimal("1"), reason="мимо API"))
        with pytest.raises(ClosedPeriodError):
            db_session.commit()
        db_session.rollback()


# ── Распределение по юрлицам: версии с 1-го числа ─────────────────────────────

class TestShareVersions:
    def test_shares_only_from_first_of_month(self, client, setup):
        h = _auth(client, setup["admin"])
        r = client.put(f"/api/employees/{setup['emp'].id}/company-shares", headers=h, json={
            "shares": [{"company_id": setup["kft"].id, "percent": "100"}],
            "effective_from": "2026-05-15",
        })
        assert r.status_code == 422

    def test_new_set_applies_from_its_month_only(self, client, db_session, setup):
        h = _auth(client, setup["admin"])
        r = client.put(f"/api/employees/{setup['emp'].id}/company-shares", headers=h, json={
            "shares": [{"company_id": setup["kft"].id, "percent": "100"}],
            "effective_from": "2026-06-01",
        })
        assert r.status_code == 200, r.text
        assert r.json()["effective_from"] == "2026-06-01"
        emp = setup["emp"]
        may = _statement(db_session, emp).rows[0]
        assert may.distribution_source == "hours"  # в мае набора ещё нет
        june = build_payroll_statement(db_session, [emp], [], 2026, 6).rows[0]
        assert june.distribution_source == "employee"
        assert [d.company_id for d in june.distribution] == [setup["kft"].id]


# ── Займ: удержания закрытых месяцев — факт из снимка ─────────────────────────

class TestLoanFacts:
    def test_changing_loan_terms_keeps_closed_deduction(self, db_session, setup):
        emp = setup["emp"]
        emp.loan_amount, emp.loan_term_months, emp.loan_start_date = (
            Decimal("12000"), 12, date(2026, 5, 1))
        db_session.commit()
        close_period(db_session, _period(db_session, setup["dept"]), setup["admin"])
        # Займ перезаняли: 6000 на 3 месяца — план 2000 в месяц.
        emp.loan_amount, emp.loan_term_months = Decimal("6000"), 3
        # В июне есть начисление: месяц без него заём пропускает (п.5.1).
        db_session.add(EmployeeAdjustment(
            employee_id=emp.id, position_id=emp.primary_position.id, year=2026, month=6,
            kind="premium", amount=Decimal("10000"), reason="июнь", created_by_id=setup["admin"].id,
        ))
        db_session.commit()
        may = build_payroll_summary(db_session, [emp], [], 2026, 5).employees[0]
        assert may.loan_deduction == Decimal("1000")  # из снимка
        june = build_payroll_summary(db_session, [emp], [], 2026, 6).employees[0]
        # Остаток июня = 6000 − 1000 фактически удержанных в мае.
        assert june.loan_deduction == Decimal("2000")
        assert june.loan_remaining == Decimal("3000")
